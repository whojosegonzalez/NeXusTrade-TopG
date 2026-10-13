import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.05
        
        # We put code in a single-cell shaded table for clean background
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.9)
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F4F6F9")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.0
        run = cp.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 45, 60)
        code_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        # Parse table lines
        header_row = [c.strip() for c in table_lines[0].split('|')[1:-1]]
        # Skip separator row (table_lines[1])
        data_rows = []
        for l in table_lines[2:]:
            cols = [c.strip() for c in l.split('|')[1:-1]]
            if cols:
                data_rows.append(cols)

        num_cols = len(header_row)
        tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, "CBD5E1")

        # Header formatting
        for col_idx, col_name in enumerate(header_row):
            cell = tbl.cell(0, col_idx)
            set_cell_background(cell, "1E293B") # Dark slate
            set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(col_name)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data formatting
        for row_idx, row_data in enumerate(data_rows):
            fill_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_data):
                if col_idx < num_cols:
                    cell = tbl.cell(row_idx + 1, col_idx)
                    set_cell_background(cell, fill_color)
                    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    add_formatted_text(p, val, font_size=Pt(9))

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        table_lines = []

    def add_formatted_text(paragraph, text, font_size=Pt(10.5), default_bold=False):
        # Clean latex math notations like \ge, \le, \text{}, etc. for docx
        cleaned = text.replace(r'\ge', '≥').replace(r'\le', '≤').replace(r'\ne', '≠')
        cleaned = cleaned.replace(r'\times', '×').replace(r'\rightarrow', '→')
        cleaned = re.sub(r'\\text\{([^}]+)\}', r'\1', cleaned)
        cleaned = cleaned.replace('$', '')

        # Match bold **text** or `code` or normal
        parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', cleaned)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
                run.font.size = font_size
            elif part.startswith('`') and part.endswith('`'):
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(font_size.pt - 1)
                run.font.color.rgb = RGBColor(180, 40, 40)
            else:
                run = paragraph.add_run(part)
                run.bold = default_bold
                run.font.size = font_size

    for line in lines:
        stripped = line.strip()

        # Code block toggle
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Table row
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table()

        # Empty lines
        if not stripped:
            continue

        # Dividers
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 23, 42)
            continue
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(30, 41, 59)
            continue
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(51, 65, 85)
            continue
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Bullet lists
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_text(p, stripped[2:])
            continue

        # Numbered lists
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_text(p, m_num.group(2))
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_text(p, stripped)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)
    print(f"✅ Generated Word document: {docx_path}")

def docx_to_pdf(docx_path, pdf_path):
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        abs_docx = os.path.abspath(docx_path)
        abs_pdf = os.path.abspath(pdf_path)
        doc = word.Documents.Open(abs_docx)
        # 17 = wdFormatPDF
        doc.SaveAs2(abs_pdf, FileFormat=17)
        doc.Close()
        print(f"✅ Generated PDF document: {pdf_path}")
    finally:
        word.Quit()

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    if len(sys.argv) >= 4:
        md_file = os.path.abspath(sys.argv[1])
        docx_file = os.path.abspath(sys.argv[2])
        pdf_file = os.path.abspath(sys.argv[3])
    else:
        md_file = os.path.join(base_dir, "nexus_trading_strategy_set.md")
        docx_file = os.path.join(base_dir, "nexus_trading_strategy_set.docx")
        pdf_file = os.path.join(base_dir, "nexus_trading_strategy_set.pdf")

    print(f"Reading markdown from: {md_file}")
    markdown_to_docx(md_file, docx_file)
    docx_to_pdf(docx_file, pdf_file)
    print("🎉 Export complete!")
